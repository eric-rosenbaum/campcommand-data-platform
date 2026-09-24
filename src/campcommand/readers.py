import csv
import io
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from campcommand.clients import SourceFormat


@dataclass
class SourceFile:
    path: Path
    headers: list[str]
    rows: list[dict[str, object]]
    row_numbers: list[int]


class SourceReadError(ValueError):
    pass


def read_source(path: Path, fmt: SourceFormat) -> SourceFile:
    if fmt.type == "csv":
        return _read_csv(path, fmt)
    if fmt.type == "excel":
        return _read_excel(path, fmt)
    raise SourceReadError(f"unsupported format {fmt.type!r}")


def _read_csv(path: Path, fmt: SourceFormat) -> SourceFile:
    try:
        text = path.read_text(encoding=fmt.encoding)
    except UnicodeDecodeError as exc:
        raise SourceReadError(f"file isn't {fmt.encoding} encoded ({exc.reason})") from None
    reader = csv.reader(io.StringIO(text.removeprefix("\ufeff"), newline=""), delimiter=fmt.delimiter)
    header_line = None
    rows, row_numbers = [], []
    # Numbered by record, which is what the client sees as the row number when
    # they open the file in Excel.
    for line_number, cells in enumerate(reader, start=1):
        if not any(cell.strip() for cell in cells):
            continue
        if header_line is None:
            if line_number < fmt.header_row:
                continue
            header_line = [cell.strip() for cell in cells]
            continue
        rows.append(dict(zip(header_line, cells, strict=False)))
        row_numbers.append(line_number)
    if header_line is None:
        raise SourceReadError("file is empty")
    return SourceFile(path, header_line, rows, row_numbers)


def _read_excel(path: Path, fmt: SourceFormat) -> SourceFile:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if fmt.sheet and fmt.sheet not in workbook.sheetnames:
            raise SourceReadError(f"no sheet named {fmt.sheet!r} (found: {', '.join(workbook.sheetnames)})")
        sheet = workbook[fmt.sheet] if fmt.sheet else workbook.worksheets[0]
        headers: list[str] | None = None
        rows, row_numbers = [], []
        for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            if row_number < fmt.header_row:
                continue
            if row_number == fmt.header_row:
                headers = [str(v).strip() if v is not None else "" for v in values]
                continue
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
                continue
            rows.append(dict(zip(headers, values, strict=False)))
            row_numbers.append(row_number)
    finally:
        workbook.close()
    if headers is None:
        raise SourceReadError(f"header row {fmt.header_row} not found")
    return SourceFile(path, headers, rows, row_numbers)
